"""按文件类型选择解析器，将上传文件转换为知识库纯文本。"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from io import BytesIO, StringIO
import json
from pathlib import Path
from typing import Callable
from urllib.parse import quote

from bs4 import BeautifulSoup
import httpx


class DocumentParseError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    parser: str


Parser = Callable[[bytes], str]


class ApacheTikaParser:
    """Extract text through an Apache Tika Server without starting a JVM in API."""

    def __init__(
        self,
        base_url: str,
        *,
        timeout: float = 120.0,
        max_extracted_chars: int = 2_000_000,
        transport: httpx.BaseTransport | None = None,
    ) -> None:
        self.base_url = base_url.rstrip("/")
        self.timeout = timeout
        self.max_extracted_chars = max_extracted_chars
        self.transport = transport

    def parse(self, filename: str, content: bytes) -> str:
        safe_name = quote(Path(filename).name, safe="._-")
        headers = {
            "Accept": "text/plain",
            "Content-Type": "application/octet-stream",
            "Content-Disposition": f"attachment; filename*=UTF-8''{safe_name}",
        }
        try:
            with httpx.Client(
                timeout=self.timeout,
                transport=self.transport,
            ) as client:
                response = client.put(
                    f"{self.base_url}/tika", content=content, headers=headers
                )
                response.raise_for_status()
        except httpx.HTTPError as exc:
            raise DocumentParseError(f"Apache Tika 解析失败：{exc}") from exc
        text = response.text.strip()
        if len(text) > self.max_extracted_chars:
            raise DocumentParseError("Apache Tika 提取文本超过长度限制")
        if not text:
            raise DocumentParseError("Apache Tika 未提取到文本")
        return text


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("无法识别文本编码")


def _parse_json(content: bytes) -> str:
    payload = json.loads(_decode_text(content))
    return json.dumps(payload, ensure_ascii=False, indent=2)


def _parse_csv(content: bytes) -> str:
    rows = csv.reader(StringIO(_decode_text(content)))
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)


def _parse_html(content: bytes) -> str:
    soup = BeautifulSoup(_decode_text(content), "lxml")
    for node in soup(["script", "style", "noscript"]):
        node.decompose()
    return soup.get_text("\n", strip=True)


def _parse_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentParseError("PDF 解析依赖 pypdf 未安装") from exc
    return "\n\n".join(
        page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages
    )


def _parse_docx(content: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentParseError("DOCX 解析依赖 python-docx 未安装") from exc
    document = Document(BytesIO(content))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _parse_pptx(content: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise DocumentParseError("PPTX 解析依赖 python-pptx 未安装") from exc
    presentation = Presentation(BytesIO(content))
    return "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text") and shape.text.strip()
    )


def _parse_xlsx(content: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DocumentParseError("XLSX 解析依赖 openpyxl 未安装") from exc
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            lines.append(
                " | ".join("" if value is None else str(value) for value in row)
            )
    return "\n".join(lines)


class DocumentParserRegistry:
    TIKA_ONLY_SUFFIXES = {
        ".ppt",
        ".rtf",
        ".odt",
        ".ods",
        ".odp",
        ".epub",
    }

    def __init__(self, tika: ApacheTikaParser | None = None) -> None:
        self.tika = tika
        self._parsers: dict[str, tuple[str, Parser]] = {}
        self.register({".txt", ".md", ".log"}, "text", _decode_text)
        self.register({".json"}, "json", _parse_json)
        self.register({".csv"}, "csv", _parse_csv)
        self.register({".html", ".htm"}, "html", _parse_html)
        self.register({".pdf"}, "pypdf", _parse_pdf)
        self.register({".docx"}, "python-docx", _parse_docx)
        self.register({".pptx"}, "python-pptx", _parse_pptx)
        self.register({".xlsx"}, "openpyxl", _parse_xlsx)

    def register(self, suffixes: set[str], name: str, parser: Parser) -> None:
        for suffix in suffixes:
            self._parsers[suffix.casefold()] = (name, parser)

    def parse(self, filename: str, content: bytes) -> ParsedDocument:
        suffix = Path(filename).suffix.casefold()
        registered = self._parsers.get(suffix)
        if registered is None:
            if self.tika is not None and suffix in self.TIKA_ONLY_SUFFIXES:
                return ParsedDocument(
                    text=self.tika.parse(filename, content),
                    parser="apache-tika",
                )
            raise DocumentParseError(f"暂不支持解析 {suffix or '无扩展名'} 文件")
        name, parser = registered
        try:
            text = parser(content).strip()
        except Exception as exc:
            if self.tika is not None:
                return ParsedDocument(
                    text=self.tika.parse(filename, content),
                    parser="apache-tika",
                )
            if isinstance(exc, DocumentParseError):
                raise
            raise DocumentParseError(f"{name} 解析失败：{exc}") from exc
        if not text:
            if self.tika is not None:
                return ParsedDocument(
                    text=self.tika.parse(filename, content),
                    parser="apache-tika",
                )
            raise DocumentParseError("文档中没有可入库文本")
        return ParsedDocument(text=text, parser=name)


__all__ = [
    "ApacheTikaParser",
    "DocumentParseError",
    "DocumentParserRegistry",
    "ParsedDocument",
]
