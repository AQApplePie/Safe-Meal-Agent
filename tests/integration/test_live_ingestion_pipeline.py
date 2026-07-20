from __future__ import annotations

import asyncio
from io import BytesIO
import os
from pathlib import Path

import pytest

from SafeMealAgent.back.bootstrap.container import AppContainer
from SafeMealAgent.back.infrastructure.ingestion.connectors import (
    LocalFileConnector,
    SourceConnectorRegistry,
)
from SafeMealAgent.back.infrastructure.ingestion.documents import ApacheTikaParser
from SafeMealAgent.back.infrastructure.ingestion.documents import DocumentParserRegistry
from SafeMealAgent.back.infrastructure.ingestion.sync import JsonSyncStateStore, SourceSyncService


def test_live_tika_extracts_legacy_office_compatible_document() -> None:
    if os.getenv("RUN_TIKA_INTEGRATION") != "1":
        pytest.skip("set RUN_TIKA_INTEGRATION=1 to run Apache Tika integration")

    rtf = (
        b"{\\rtf1\\ansi\\deff0 {\\fonttbl {\\f0 Times New Roman;}}"
        b"\\f0\\fs24 SafeMeal Agent Tika legacy pipeline\\par}"
    )
    parsed = ApacheTikaParser("http://127.0.0.1:9998").parse("legacy.rtf", rtf)

    assert "SafeMeal Agent Tika legacy pipeline" in parsed


def test_live_tika_ocr_fallback_extracts_scanned_pdf() -> None:
    if os.getenv("RUN_TIKA_INTEGRATION") != "1":
        pytest.skip("set RUN_TIKA_INTEGRATION=1 to run Apache Tika integration")

    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1200, 300), "white")
    draw = ImageDraw.Draw(image)
    draw.text(
        (40, 90),
        "SAFEMEAL OCR 7319",
        fill="black",
        font=ImageFont.load_default(size=72),
    )
    scanned_pdf = BytesIO()
    image.save(scanned_pdf, format="PDF", resolution=150)
    registry = DocumentParserRegistry(
        tika=ApacheTikaParser("http://127.0.0.1:9998", timeout=180)
    )

    parsed = registry.parse("scanned.pdf", scanned_pdf.getvalue())

    assert parsed.parser == "apache-tika"
    assert "SAFEMEAL" in parsed.text.upper()


def test_live_local_docx_to_embedding_milvus_and_rerank(tmp_path: Path) -> None:
    if os.getenv("RUN_INGESTION_INTEGRATION") != "1":
        pytest.skip(
            "set RUN_INGESTION_INTEGRATION=1 to run the live ingestion pipeline"
        )

    from docx import Document

    document_id = "local:pipeline.docx"
    docx_file = BytesIO()
    document = Document()
    document.add_paragraph(
        "数据管道验收暗号是星河青椒。这个暗号只用于验证文档抓取、解析、"
        "语义分块、向量化、Milvus 入库和 Rerank 检索。"
    )
    document.add_paragraph("汽车发动机属于另一个语义主题，用于验证语义边界。")
    document.save(docx_file)
    source_file = tmp_path / "pipeline.docx"
    source_file.write_bytes(docx_file.getvalue())

    async def scenario() -> tuple[str, list[dict], str]:
        container = AppContainer()
        connectors = SourceConnectorRegistry()
        connectors.register("local", LocalFileConnector(tmp_path))
        knowledge = await container.get_knowledge_service()
        service = SourceSyncService(
            connectors,
            knowledge,
            JsonSyncStateStore(tmp_path / "states.json"),
            parsers=container.document_parsers,
        )
        try:
            synced = await service.sync("local://pipeline.docx", chunk_strategy="auto")
            results = await knowledge.search("星河青椒是什么数据管道暗号", top_k=5)
            source_file.unlink()
            deleted = await service.sync("local://pipeline.docx")
            return synced.status, results, deleted.status
        finally:
            await knowledge.delete_document(document_id)
            await container.shutdown()

    synced_status, results, deleted_status = asyncio.run(scenario())
    assert synced_status == "synced"
    assert deleted_status == "deleted"
    assert results
    assert any("星河青椒" in str(item.get("content")) for item in results)
    assert all("rerank_score" in item for item in results)
