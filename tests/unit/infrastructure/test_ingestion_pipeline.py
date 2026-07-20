from __future__ import annotations

import asyncio
from io import BytesIO
from pathlib import Path
from typing import Sequence

import httpx
import pytest

from SafeMealAgent.back.application.use_cases.knowledge.chunking import (
    ChunkStrategySelector,
    RecursiveChunkStrategy,
    SemanticChunkStrategy,
)
from SafeMealAgent.back.infrastructure.ingestion.connectors import (
    FetchedSource,
    FeishuDocumentConnector,
    HTTPSourceConnector,
    LocalFileConnector,
    S3SourceConnector,
    SourceConnectorRegistry,
    SourceNotFoundError,
)
from SafeMealAgent.back.infrastructure.ingestion.documents import (
    ApacheTikaParser,
    DocumentParserRegistry,
)
from SafeMealAgent.back.infrastructure.ingestion.sync import (
    JsonSyncStateStore,
    SourceSyncService,
    SyncJob,
    SyncScheduler,
)


class _SemanticEmbedder:
    def embed_documents(self, texts: Sequence[str]) -> list[list[float]]:
        return [
            [1.0, 0.0] if "水果" in text or "苹果" in text else [0.0, 1.0]
            for text in texts
        ]

    def embed_query(self, text: str) -> list[float]:
        return [1.0, 0.0]


def test_local_http_s3_and_feishu_connectors_fetch_content(tmp_path: Path) -> None:
    local_file = tmp_path / "local.txt"
    local_file.write_text("local content", encoding="utf-8")
    local = asyncio.run(LocalFileConnector(tmp_path).fetch("local.txt"))

    def http_handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, content=b"https content")

    http = asyncio.run(
        HTTPSourceConnector(
            {"docs.example.com"},
            transport=httpx.MockTransport(http_handler),
        ).fetch("https://docs.example.com/guide.txt")
    )

    class Body:
        def read(self, limit: int) -> bytes:
            return b"s3 content"

    class S3Client:
        def get_object(self, **kwargs: object) -> dict[str, object]:
            return {"ContentLength": 10, "Body": Body()}

    s3 = asyncio.run(
        S3SourceConnector(
            bucket="knowledge",
            prefix="docs",
            client_factory=S3Client,
        ).fetch("docs/guide.txt")
    )

    def feishu_handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            json={"code": 0, "data": {"content": "feishu content"}},
        )

    feishu = asyncio.run(
        FeishuDocumentConnector(
            "token",
            transport=httpx.MockTransport(feishu_handler),
        ).fetch("document-token")
    )

    assert local.content == b"local content"
    assert http.content == b"https content"
    assert s3.content == b"s3 content"
    assert feishu.content == b"feishu content"


def test_source_connectors_enforce_traversal_host_and_prefix_boundaries(
    tmp_path: Path,
) -> None:
    with pytest.raises(ValueError, match="escapes allowed root"):
        asyncio.run(LocalFileConnector(tmp_path).fetch("../secret.txt"))
    with pytest.raises(ValueError, match="not allowlisted"):
        asyncio.run(
            HTTPSourceConnector({"allowed.example"}).fetch(
                "https://blocked.example/file.txt"
            )
        )
    with pytest.raises(ValueError, match="outside configured prefix"):
        asyncio.run(
            S3SourceConnector(
                bucket="knowledge",
                prefix="docs",
                client_factory=lambda: object(),
            ).fetch("private/file.txt")
        )


def test_auto_chunking_selects_semantic_strategy_by_document_type() -> None:
    recursive = RecursiveChunkStrategy(chunk_size=200, chunk_overlap=0)
    semantic = SemanticChunkStrategy(
        embedder=_SemanticEmbedder(),
        recursive=recursive,
        chunk_size=200,
        chunk_overlap=0,
        breakpoint_threshold=0.5,
        max_segments=20,
    )
    selector = ChunkStrategySelector(
        default_strategy="recursive",
        document_strategies={"pypdf": "semantic"},
        recursive=recursive,
        semantic=semantic,
    )

    name, documents = selector.split(
        "苹果是一种水果。香蕉也是水果。汽车使用发动机。道路用于交通。",
        {"parser": "pypdf"},
        "auto",
    )

    assert name == "semantic"
    assert len(documents) == 2
    assert all(item.metadata["chunk_strategy"] == "semantic" for item in documents)


def test_explicit_recursive_strategy_overrides_document_mapping() -> None:
    recursive = RecursiveChunkStrategy(chunk_size=200, chunk_overlap=0)
    selector = ChunkStrategySelector(
        default_strategy="recursive",
        document_strategies={"pypdf": "semantic"},
        recursive=recursive,
        semantic=SemanticChunkStrategy(
            embedder=_SemanticEmbedder(),
            recursive=recursive,
            chunk_size=200,
            chunk_overlap=0,
            breakpoint_threshold=0.5,
            max_segments=20,
        ),
    )

    name, documents = selector.split("短文档。", {"parser": "pypdf"}, "recursive")

    assert name == "recursive"
    assert len(documents) == 1


def test_tika_parses_tika_only_formats_and_is_a_native_parser_fallback() -> None:
    requests: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requests.append(request.headers["content-disposition"])
        return httpx.Response(200, text="Tika 提取的正文")

    registry = DocumentParserRegistry(
        tika=ApacheTikaParser(
            "http://tika.test:9998",
            transport=httpx.MockTransport(handler),
        )
    )

    legacy = registry.parse("legacy.ppt", b"legacy office bytes")
    fallback = registry.parse("broken.pdf", b"not a valid PDF")

    assert legacy.parser == "apache-tika"
    assert fallback.parser == "apache-tika"
    assert len(requests) == 2
    assert "legacy.ppt" in requests[0]


def test_docx_and_pptx_native_parsers_are_runnable() -> None:
    from docx import Document
    from pptx import Presentation

    docx_file = BytesIO()
    document = Document()
    document.add_paragraph("DOCX 正文")
    document.save(docx_file)

    pptx_file = BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "PPTX 正文"
    presentation.save(pptx_file)

    registry = DocumentParserRegistry()
    assert "DOCX 正文" in registry.parse("sample.docx", docx_file.getvalue()).text
    assert "PPTX 正文" in registry.parse("sample.pptx", pptx_file.getvalue()).text


class _ConstantConnector:
    def __init__(self) -> None:
        self.missing = False

    async def fetch(self, reference: str) -> FetchedSource:
        if self.missing:
            raise SourceNotFoundError(reference)
        return FetchedSource("local:doc.txt", "doc.txt", b"incremental document")


class _KnowledgeRecorder:
    def __init__(self, *, fail_first: bool = False) -> None:
        self.fail_first = fail_first
        self.ingests = 0
        self.deleted: list[str] = []

    async def ingest_text(self, text: str, **kwargs: object) -> dict[str, object]:
        self.ingests += 1
        if self.fail_first and self.ingests == 1:
            raise RuntimeError("temporary embedding failure")
        return {"add_count": 1}

    async def delete_document(self, document_id: str) -> bool:
        self.deleted.append(document_id)
        return True


def _sync_service(
    tmp_path: Path,
    connector: _ConstantConnector,
    knowledge: _KnowledgeRecorder,
) -> SourceSyncService:
    registry = SourceConnectorRegistry()
    registry.register("local", connector)
    return SourceSyncService(
        registry,
        knowledge,  # type: ignore[arg-type]
        JsonSyncStateStore(tmp_path / "states.json"),
    )


def test_failed_document_is_retried_even_when_checksum_is_unchanged(
    tmp_path: Path,
) -> None:
    connector = _ConstantConnector()
    knowledge = _KnowledgeRecorder(fail_first=True)
    service = _sync_service(tmp_path, connector, knowledge)

    first = asyncio.run(service.sync("local://doc.txt"))
    second = asyncio.run(service.sync("local://doc.txt"))

    assert first.status == "failed"
    assert second.status == "synced"
    assert knowledge.ingests == 2


def test_missing_source_deletes_previous_vector_document(tmp_path: Path) -> None:
    connector = _ConstantConnector()
    knowledge = _KnowledgeRecorder()
    service = _sync_service(tmp_path, connector, knowledge)

    first = asyncio.run(service.sync("local://doc.txt"))
    connector.missing = True
    deleted = asyncio.run(service.sync("local://doc.txt"))

    assert first.status == "synced"
    assert deleted.status == "deleted"
    assert knowledge.deleted == ["local:doc.txt"]


def test_scheduler_isolates_one_iteration_failure_and_keeps_running() -> None:
    class FlakyService:
        def __init__(self) -> None:
            self.calls = 0
            self.completed = asyncio.Event()

        async def sync(self, source_uri: str, **kwargs: object) -> object:
            self.calls += 1
            if self.calls == 1:
                raise RuntimeError("temporary connector failure")
            self.completed.set()
            return object()

    async def scenario() -> int:
        service = FlakyService()
        scheduler = SyncScheduler(
            service,  # type: ignore[arg-type]
            [SyncJob("local://doc.txt", interval_seconds=0.01)],
        )
        scheduler.start()
        await asyncio.wait_for(service.completed.wait(), timeout=1)
        await scheduler.stop()
        return service.calls

    assert asyncio.run(scenario()) >= 2
